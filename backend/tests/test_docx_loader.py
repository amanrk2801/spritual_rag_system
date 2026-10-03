import docx
from docx.enum.text import WD_BREAK

from app.ingestion.docx_loader import load_docx


def test_docx_headings_merged_and_pages_tracked(tmp_path):
    d = docx.Document()
    d.add_paragraph("CHAPTER I")
    d.add_paragraph("INTRODUCTORY")
    d.add_paragraph("Each soul is potentially divine. This is a normal paragraph.")
    d.add_paragraph("Before the break.").add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("PRANA", level=2)
    d.add_paragraph("Pranayama is the control of prana.")
    path = tmp_path / "book.docx"
    d.save(path)

    pages = dict(load_docx(path))
    assert pages[1].startswith("## CHAPTER I — INTRODUCTORY")
    assert "Before the break." in pages[1]
    assert pages[2].startswith("## PRANA") and "control of prana" in pages[2]
